#!/usr/bin/env python3
"""
PtolC/tools/ingest_system.py — full filesystem ingest for ptolemy (Python3)

Recursively visits every readable file from ROOT, extracts text,
feeds it to ptolemy -l - in batches.  State saved to JSON every
BATCH_DIRS directories — safe to kill and resume at any time.

SETUP (run once):
    cd PtolC
    make
    bash tools/install_ingest_deps.sh
    ./tools/checkpoint_expand monad_wordnet.bin 255400
    ./ptolemy -s

RUN:
    python3 tools/ingest_system.py --root / --first /SystemTree.txt
    # overnight:
    nohup python3 tools/ingest_system.py --root / --first /SystemTree.txt \\
        2>&1 | tee ingest_py.log &

RESUME (after kill):
    python3 tools/ingest_system.py --root / --first /SystemTree.txt
    # Already-completed directories are skipped automatically.

ANDROID FACTORY IMAGES / PARTITION DUMPS:
    Mount the system.img before running:
        sudo mount -o loop,ro system.img /mnt/android_system
    Then pass as an extra root:
        python3 tools/ingest_system.py --root / --extra /mnt/android_system
    Or without root mount, use e2tools:
        e2ls system.img:/  (read-only, no mount needed)
    ADB shell (rooted phone):
        adb shell find / -readable -type f 2>/dev/null \\
            | while read f; do adb shell cat "$f" 2>/dev/null; done \\
            | ./ptolemy -l -

GOOGLE DRIVE (GNOME GVFS):
    Detected automatically at /run/user/<uid>/gvfs/
    Mount it in Nautilus (Files) before running — it appears as
    google-drive:host=gmail.com,user=... subdirectory.
"""

import argparse
import json
import logging
import math
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

# ── Optional library imports ───────────────────────────────────────────────────
try:
    import magic as libmagic
    _HAVE_MAGIC = True
except ImportError:
    _HAVE_MAGIC = False

try:
    import chardet
    _HAVE_CHARDET = True
except ImportError:
    _HAVE_CHARDET = False

try:
    from pdfminer.high_level import extract_text as pdf_extract
    _HAVE_PDFMINER = True
except ImportError:
    _HAVE_PDFMINER = False

try:
    import docx as python_docx
    _HAVE_DOCX = True
except ImportError:
    _HAVE_DOCX = False

try:
    import ebooklib
    from ebooklib import epub as epub_lib
    import re as _re
    _HAVE_EPUB = True
except ImportError:
    _HAVE_EPUB = False

try:
    from odf import text as odf_text, teletype
    from odf.opendocument import load as odf_load
    _HAVE_ODF = True
except ImportError:
    _HAVE_ODF = False

try:
    from striprtf.striprtf import rtf_to_text
    _HAVE_RTF = True
except ImportError:
    _HAVE_RTF = False

try:
    from bs4 import BeautifulSoup
    _HAVE_BS4 = True
except ImportError:
    _HAVE_BS4 = False

try:
    import pytesseract
    from PIL import Image
    _HAVE_OCR = True
except ImportError:
    _HAVE_OCR = False

# ── Constants ──────────────────────────────────────────────────────────────────
CHUNK_LINES   = 50_000       # max lines per ptolemy -l - call (large file slicing)
DEFAULT_BATCH = 10           # checkpoint every N directories
STATE_FILE    = "ingest_state.json"
PTOLEMY       = "./ptolemy"

PRUNE_DIRS = frozenset([
    "/proc", "/sys", "/dev", "/run/lock",
    "/snap", "/var/run", "/tmp",
])

LOG = logging.getLogger("ingest")

# ── Mime detection ─────────────────────────────────────────────────────────────

def mime_type(path: str) -> str:
    if _HAVE_MAGIC:
        try:
            return libmagic.from_file(path, mime=True) or ""
        except Exception:
            pass
    result = subprocess.run(
        ["file", "-b", "--mime-type", path],
        capture_output=True, text=True, timeout=5
    )
    return result.stdout.strip() if result.returncode == 0 else ""

# ── Text extraction ────────────────────────────────────────────────────────────

def decode_bytes(raw: bytes) -> str:
    if _HAVE_CHARDET:
        det = chardet.detect(raw[:8192])
        enc = det.get("encoding") or "utf-8"
    else:
        enc = "utf-8"
    return raw.decode(enc, errors="replace")


def extract_text(path: str, mime: str) -> str:
    """Return extracted text from path (may be empty string on failure)."""
    try:
        if mime.startswith("text/"):
            with open(path, "rb") as f:
                return decode_bytes(f.read())

        if mime == "application/pdf":
            if _HAVE_PDFMINER:
                try:
                    return pdf_extract(path) or ""
                except Exception:
                    pass
            result = subprocess.run(
                ["pdftotext", "-q", "-nopgbrk", "-enc", "UTF-8", path, "-"],
                capture_output=True, timeout=60
            )
            if result.returncode == 0:
                return result.stdout.decode("utf-8", errors="replace")
            # OCR fallback for image-only PDFs
            if _HAVE_OCR:
                try:
                    import pdf2image
                    imgs = pdf2image.convert_from_path(path, dpi=150)
                    return "\n".join(pytesseract.image_to_string(img) for img in imgs)
                except Exception:
                    pass
            return ""

        if mime in ("application/gzip", "application/x-gzip"):
            result = subprocess.run(
                ["zcat", path], capture_output=True, timeout=30
            )
            if result.returncode != 0:
                return ""
            inner = subprocess.run(
                ["file", "-b", "--mime-type", "-"],
                input=result.stdout[:512], capture_output=True, timeout=5
            ).stdout.strip().decode()
            if inner.startswith("text/"):
                col = subprocess.run(
                    ["col", "-b"], input=result.stdout,
                    capture_output=True, timeout=30
                )
                return col.stdout.decode("utf-8", errors="replace")
            return ""

        if mime in ("application/x-bzip2",):
            result = subprocess.run(
                ["bzcat", path], capture_output=True, timeout=30
            )
            return result.stdout.decode("utf-8", errors="replace") if result.returncode == 0 else ""

        if mime in ("application/x-xz",):
            result = subprocess.run(
                ["xzcat", path], capture_output=True, timeout=30
            )
            return result.stdout.decode("utf-8", errors="replace") if result.returncode == 0 else ""

        if mime in ("text/html", "application/xhtml+xml"):
            with open(path, "rb") as f:
                raw = f.read()
            if _HAVE_BS4:
                soup = BeautifulSoup(raw, "lxml")
                return soup.get_text(separator="\n")
            result = subprocess.run(
                ["html2text", path], capture_output=True, timeout=30
            )
            return result.stdout.decode("utf-8", errors="replace") if result.returncode == 0 else ""

        if mime == "application/msword":
            result = subprocess.run(
                ["catdoc", path], capture_output=True, timeout=30
            )
            return result.stdout.decode("utf-8", errors="replace") if result.returncode == 0 else ""

        if mime in ("application/rtf", "text/rtf"):
            if _HAVE_RTF:
                with open(path, "rb") as f:
                    return rtf_to_text(f.read().decode("utf-8", errors="replace"))
            result = subprocess.run(
                ["unrtf", "--text", path], capture_output=True, timeout=30
            )
            return result.stdout.decode("utf-8", errors="replace") if result.returncode == 0 else ""

        if mime == "application/epub+zip":
            if _HAVE_EPUB:
                try:
                    book = epub_lib.read_epub(path, options={"ignore_ncx": True})
                    parts = []
                    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                        soup = BeautifulSoup(item.get_content(), "lxml") if _HAVE_BS4 else None
                        if soup:
                            parts.append(soup.get_text(separator="\n"))
                        else:
                            parts.append(_re.sub(r"<[^>]+>", "", item.get_content().decode("utf-8", errors="replace")))
                    return "\n".join(parts)
                except Exception:
                    pass
            import zipfile, re
            try:
                with zipfile.ZipFile(path) as z:
                    parts = []
                    for n in z.namelist():
                        if n.endswith((".xhtml", ".html", ".htm")):
                            raw = z.read(n).decode("utf-8", errors="replace")
                            parts.append(re.sub(r"<[^>]+>", "", raw))
                    return "\n".join(parts)
            except Exception:
                return ""

        if mime == "application/vnd.oasis.opendocument.text":
            if _HAVE_ODF:
                try:
                    doc = odf_load(path)
                    return teletype.extractText(doc.text)
                except Exception:
                    pass
            result = subprocess.run(
                ["odt2txt", "--stdout", path], capture_output=True, timeout=30
            )
            return result.stdout.decode("utf-8", errors="replace") if result.returncode == 0 else ""

        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            if _HAVE_DOCX:
                try:
                    doc = python_docx.Document(path)
                    return "\n".join(p.text for p in doc.paragraphs)
                except Exception:
                    pass
            return ""

        if mime in ("application/vnd.ms-excel",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                rows = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        rows.append("\t".join(str(c) for c in row if c is not None))
                return "\n".join(rows)
            except Exception:
                return ""

        if mime.startswith("image/") and _HAVE_OCR:
            try:
                img = Image.open(path)
                return pytesseract.image_to_string(img)
            except Exception:
                pass

        # .po / gettext — parse msgstr lines (Unicode locale data)
        if path.endswith(".po") or path.endswith(".pot"):
            try:
                with open(path, "rb") as f:
                    raw = f.read()
                text = decode_bytes(raw)
                lines = [l[7:].strip().strip('"') for l in text.splitlines()
                         if l.startswith("msgstr") or l.startswith("msgid")]
                return "\n".join(l for l in lines if l)
            except Exception:
                return ""

    except (PermissionError, FileNotFoundError, IsADirectoryError, OSError):
        pass
    except subprocess.TimeoutExpired:
        LOG.warning("timeout: %s", path)
    except Exception as e:
        LOG.debug("extract error %s: %s", path, e)

    return ""

# ── Ptolemy pipe ───────────────────────────────────────────────────────────────

def learn_text(text: str, ptolemy: str) -> bool:
    """Feed text to ptolemy -l -, return True on success."""
    if not text.strip():
        return True
    try:
        result = subprocess.run(
            [ptolemy, "-l", "-"],
            input=text.encode("utf-8", errors="replace"),
            capture_output=True, timeout=300
        )
        return result.returncode == 0
    except subprocess.TimeoutExpired:
        LOG.warning("ptolemy timeout during learn")
        return False
    except Exception as e:
        LOG.error("ptolemy error: %s", e)
        return False


def learn_chunked(text: str, ptolemy: str, chunk_lines: int = CHUNK_LINES) -> int:
    """Learn text in chunks of chunk_lines lines. Returns number of chunks sent."""
    lines = text.splitlines(keepends=True)
    if not lines:
        return 0
    chunks = math.ceil(len(lines) / chunk_lines)
    for i in range(chunks):
        chunk = "".join(lines[i * chunk_lines:(i + 1) * chunk_lines])
        if not learn_text(chunk, ptolemy):
            LOG.warning("learn failed on chunk %d/%d", i + 1, chunks)
    return chunks

# ── State management ───────────────────────────────────────────────────────────

def load_state(state_path: str) -> dict:
    if os.path.exists(state_path):
        try:
            with open(state_path) as f:
                return json.load(f)
        except Exception:
            pass
    return {"done_dirs": [], "files_total": 0, "dirs_total": 0,
            "errors": 0, "started": time.time()}


def save_state(state: dict, state_path: str):
    tmp = state_path + ".tmp"
    with open(tmp, "w") as f:
        json.dump(state, f, indent=2)
    os.replace(tmp, state_path)

# ── Directory discovery ────────────────────────────────────────────────────────

def extra_roots(uid: int, user: str) -> list[str]:
    """Detect GVFS mounts and external drives."""
    extras = []

    gvfs = f"/run/user/{uid}/gvfs"
    if os.path.isdir(gvfs):
        for entry in os.scandir(gvfs):
            if entry.is_dir():
                LOG.info("GVFS mount detected: %s", entry.path)
                extras.append(entry.path)

    media = f"/media/{user}"
    if os.path.isdir(media):
        for entry in os.scandir(media):
            if entry.is_dir():
                LOG.info("External drive detected: %s", entry.path)
                extras.append(entry.path)

    return extras


def should_prune(path: str) -> bool:
    return any(path == p or path.startswith(p + "/") for p in PRUNE_DIRS)

# ── Priority file ingestion ────────────────────────────────────────────────────

def ingest_priority_file(path: str, ptolemy: str, chunk_lines: int):
    if not os.path.isfile(path):
        LOG.info("priority file not found yet: %s (will ingest when available)", path)
        return
    size = os.path.getsize(path)
    LOG.info("priority file: %s  (%.1f MB)", path, size / 1e6)
    try:
        with open(path, "rb") as f:
            raw = f.read()
        text = decode_bytes(raw)
        chunks = learn_chunked(text, ptolemy, chunk_lines)
        LOG.info("priority file done: %d chunks", chunks)
    except Exception as e:
        LOG.error("priority file error: %s", e)

# ── Main ingest loop ───────────────────────────────────────────────────────────

def ingest_tree(root: str, ptolemy: str, state: dict, state_path: str,
                batch_dirs: int, chunk_lines: int):
    done_set = set(state["done_dirs"])
    dirs_since_save = 0
    accumulated = []
    accumulated_lines = 0

    def flush():
        nonlocal accumulated, accumulated_lines
        if accumulated:
            text = "\n".join(accumulated)
            learn_chunked(text, ptolemy, chunk_lines)
            accumulated = []
            accumulated_lines = 0

    for dirpath, dirnames, filenames in os.walk(root, topdown=True, followlinks=False):
        # Prune in-place so os.walk doesn't descend
        dirnames[:] = [
            d for d in sorted(dirnames)
            if not should_prune(os.path.join(dirpath, d))
            and not os.path.islink(os.path.join(dirpath, d))
        ]

        if should_prune(dirpath):
            continue

        if dirpath in done_set:
            LOG.debug("skip (done): %s", dirpath)
            continue

        LOG.info("[%s] %s  (files=%d)",
                 time.strftime("%H:%M:%S"), dirpath, len(filenames))

        for fname in sorted(filenames):
            fpath = os.path.join(dirpath, fname)
            if os.path.islink(fpath):
                continue
            try:
                stat = os.stat(fpath)
                if stat.st_size == 0 or stat.st_size > 512 * 1024 * 1024:
                    continue
            except OSError:
                continue

            try:
                mime = mime_type(fpath)
            except Exception:
                mime = ""

            text = extract_text(fpath, mime)
            if text.strip():
                lines = text.splitlines()
                accumulated.extend(lines)
                accumulated_lines += len(lines)
                state["files_total"] += 1

                # Flush if accumulated buffer is large enough
                if accumulated_lines >= chunk_lines:
                    flush()

        # Mark directory done and possibly save state
        state["done_dirs"].append(dirpath)
        done_set.add(dirpath)
        state["dirs_total"] += 1
        dirs_since_save += 1

        if dirs_since_save >= batch_dirs:
            flush()
            save_state(state, state_path)
            elapsed = time.time() - state["started"]
            LOG.info("checkpoint: %d dirs  %d files  %.0fs elapsed",
                     state["dirs_total"], state["files_total"], elapsed)
            dirs_since_save = 0

    flush()
    save_state(state, state_path)

# ── Entry point ────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(
        description="ptolemy full-system ingest (Python3, resumable)"
    )
    ap.add_argument("--root",       default="/",          help="root directory to traverse")
    ap.add_argument("--first",      default="",           help="file to ingest first (e.g. /SystemTree.txt)")
    ap.add_argument("--ptolemy",    default=PTOLEMY,      help="path to ptolemy binary")
    ap.add_argument("--state",      default=STATE_FILE,   help="JSON state file path")
    ap.add_argument("--batch-dirs", type=int, default=DEFAULT_BATCH, help="checkpoint every N dirs")
    ap.add_argument("--chunk-lines",type=int, default=CHUNK_LINES,   help="max lines per ptolemy call")
    ap.add_argument("--extra",      action="append", default=[],      help="extra root (repeat for multiple)")
    ap.add_argument("--no-auto-extra", action="store_true",           help="skip GVFS/external drive detection")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s  %(levelname)-5s  %(message)s",
        datefmt="%H:%M:%S",
    )

    if not os.path.isfile(args.ptolemy) or not os.access(args.ptolemy, os.X_OK):
        LOG.error("ptolemy not found or not executable: %s", args.ptolemy)
        LOG.error("Build with: cd PtolC && make")
        sys.exit(1)

    # Quick sanity check
    r = subprocess.run([args.ptolemy, "-s"], capture_output=True, timeout=10)
    if r.returncode != 0:
        LOG.warning("ptolemy -s returned non-zero — continuing anyway")
    else:
        LOG.info("ptolemy status:\n%s", r.stdout.decode("utf-8", errors="replace").strip())

    state = load_state(args.state)
    if state["dirs_total"] > 0:
        LOG.info("RESUMING: %d dirs and %d files already done",
                 state["dirs_total"], state["files_total"])

    # ── Priority file first ───────────────────────────────────────────────────
    if args.first:
        if args.first not in state.get("done_dirs", []):
            LOG.info("=== Priority file ===")
            ingest_priority_file(args.first, args.ptolemy, args.chunk_lines)
            state["done_dirs"].append(args.first)
            save_state(state, args.state)

    # ── Auto-detect GVFS / external drives ───────────────────────────────────
    uid  = os.getuid()
    user = os.environ.get("USER", os.path.basename(os.path.expanduser("~")))
    extra_roots_list = list(args.extra)
    if not args.no_auto_extra:
        extra_roots_list += extra_roots(uid, user)

    # ── Main filesystem traversal ─────────────────────────────────────────────
    all_roots = [args.root] + extra_roots_list
    for root in all_roots:
        if not os.path.isdir(root):
            LOG.warning("root not accessible: %s", root)
            continue
        LOG.info("=== Traversing: %s ===", root)
        ingest_tree(root, args.ptolemy, state, args.state,
                    args.batch_dirs, args.chunk_lines)

    # ── Final report ──────────────────────────────────────────────────────────
    elapsed = time.time() - state["started"]
    LOG.info("════════════════════════════════════════")
    LOG.info("INGEST COMPLETE")
    LOG.info("  dirs  visited : %d", state["dirs_total"])
    LOG.info("  files learned : %d", state["files_total"])
    LOG.info("  elapsed       : %.0f s  (%.1f h)", elapsed, elapsed / 3600)
    LOG.info("════════════════════════════════════════")

    r = subprocess.run([args.ptolemy, "-s"], capture_output=True, timeout=10)
    if r.returncode == 0:
        print(r.stdout.decode("utf-8", errors="replace"))


if __name__ == "__main__":
    main()
